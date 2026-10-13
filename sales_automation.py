
import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# 1. Get input filename
file_name = input("Enter your CSV filename: ").strip()

if not os.path.isfile(file_name):
    print("Error: File not found!")

else:
    try:
        # 2. Load CSV
        df = pd.read_csv(file_name)
        df.columns = df.columns.str.strip()
        print("\nCSV file loaded successfully!")

        # 3. Validate required columns
        required_columns = ["Name", "Product", "Amount", "City"]
        missing_columns = [
            col for col in required_columns if col not in df.columns
        ]

        if missing_columns:
            print("Error: Required columns are missing:", missing_columns)

        else:
            # 4. Validate Amount values
            amount_text = df["Amount"].astype("string").str.strip()
            nonblank = amount_text.notna() & amount_text.ne("")
            numeric_amount = pd.to_numeric(
                amount_text, errors="coerce"
            )
            invalid_amounts = nonblank & numeric_amount.isna()

            if invalid_amounts.any():
                print("\nError: Invalid amounts found!")

                error_report = df.loc[invalid_amounts].copy()
                error_report["Error"] = "Invalid Amount"

                error_report.to_csv("error_report.csv", index=False)

                print(error_report[["Name", "Amount", "Error"]])
                print("\nError report saved to error_report.csv")
                print("Please correct the invalid amounts and run again.")

            else:
                # 5. Clean whitespace
                df["Amount"] = numeric_amount

                for col in df.select_dtypes(include=["object"]).columns:
                    df[col] = df[col].str.strip()

                df = df.replace(r"^\s*$", pd.NA, regex=True)

                # 6. Remove duplicate records
                original_rows = len(df)
                df = df.drop_duplicates().reset_index(drop=True)
                duplicates_removed = original_rows - len(df)

                print("\n--- DATA CLEANING SUMMARY ---")
                print("Rows before duplicate removal:", original_rows)
                print("Duplicates removed:", duplicates_removed)
                print("Rows after cleaning:", len(df))
                print("\nMissing values:")
                print(df.isnull().sum())

                # 7. Save cleaned data
                df.to_csv("cleaned_sales.csv", index=False)

                # 8. Calculate sales statistics
                total_sales = df["Amount"].sum()
                average_sale = df["Amount"].mean()
                highest_sale = df["Amount"].max()
                lowest_sale = df["Amount"].min()

                sales_by_city = df.groupby(
                    "City", dropna=False
                )["Amount"].sum(min_count=1)

                sales_by_product = df.groupby(
                    "Product", dropna=False
                )["Amount"].sum(min_count=1)

                print("\n--- SALES REPORT ---")
                print("Total sales:", total_sales)
                print("Average sale:", average_sale)
                print("Highest sale:", highest_sale)
                print("Lowest sale:", lowest_sale)
                print("\nSales by city:")
                print(sales_by_city)
                print("\nSales by product:")
                print(sales_by_product)

                # 9. Create report tables
                summary = pd.DataFrame({
                    "Metric": [
                        "Total Sales",
                        "Average Sale",
                        "Highest Sale",
                        "Lowest Sale"
                    ],
                    "Value": [
                        total_sales,
                        average_sale,
                        highest_sale,
                        lowest_sale
                    ]
                })

                summary.to_csv("sales_report.csv", index=False)

                sales_by_city.rename("Total Sales").rename_axis(
                    "City"
                ).reset_index().to_csv(
                    "sales_by_city.csv", index=False
                )

                sales_by_product.rename("Total Sales").rename_axis(
                    "Product"
                ).reset_index().to_csv(
                    "sales_by_product.csv", index=False
                )

                # 10. Export Excel workbook
                with pd.ExcelWriter(
                    "sales_report.xlsx", engine="openpyxl"
                ) as writer:

                    summary.to_excel(
                        writer, sheet_name="Sales Summary", index=False
                    )

                    sales_by_city.rename("Total Sales").rename_axis(
                        "City"
                    ).reset_index().to_excel(
                        writer, sheet_name="Sales by City", index=False
                    )

                    sales_by_product.rename("Total Sales").rename_axis(
                        "Product"
                    ).reset_index().to_excel(
                        writer, sheet_name="Sales by Product", index=False
                    )

                    df.to_excel(
                        writer, sheet_name="Cleaned Data", index=False
                    )

                    for ws in writer.sheets.values():
                        # Style headers
                        for cell in ws[1]:
                            cell.font = Font(
                                bold=True, color="FFFFFF"
                            )
                            cell.fill = PatternFill(
                                "solid", fgColor="1F4E78"
                            )
                            cell.alignment = Alignment(
                                horizontal="center"
                            )

                        # Adjust column widths
                        for cells in ws.columns:
                            max_length = max(
                                len(str(cell.value))
                                if cell.value is not None else 0
                                for cell in cells
                            )
                            letter = get_column_letter(
                                cells[0].column
                            )
                            ws.column_dimensions[letter].width = min(
                                max_length + 3, 35
                            )

                        # Format numeric values
                        for row in ws.iter_rows(min_row=2):
                            for cell in row:
                                if isinstance(cell.value, (int, float)):
                                    cell.number_format = '₹#,##0.00'

                        ws.freeze_panes = "A2"

                print("\n--- EXPORT COMPLETE ---")
                print("1. cleaned_sales.csv")
                print("2. sales_report.csv")
                print("3. sales_by_city.csv")
                print("4. sales_by_product.csv")
                print("5. sales_report.xlsx")

    except pd.errors.EmptyDataError:
        print("Error: The CSV file is empty.")

    except (
        pd.errors.ParserError,
        UnicodeDecodeError,
        OSError,
        ValueError
    ) as error:
        print("Error reading or processing the CSV file:", error)
